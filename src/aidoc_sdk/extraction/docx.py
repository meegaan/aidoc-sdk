from __future__ import annotations
from docx.oxml.ns import qn
from typing import Any, List

from aidoc_sdk.extraction.base import BaseExtractor
from aidoc_sdk.extraction.registry import extractor
from aidoc_sdk.core.raw import (
    RawDocument,
    RawParagraph,
    RawHeading,
    RawTable,
)
from aidoc_sdk.exceptions import AIDocTypeError, AIDocValueError


# ---------------------------------------------------------------------------
# DOCX Extractor (v1)
# ---------------------------------------------------------------------------
# - Uses python-docx (optional dependency)
# - Extracts paragraphs, headings, tables
# - Preserves document order
# - No normalization
# - No hashing
# ---------------------------------------------------------------------------


@extractor("docx")
class DOCXExtractor(BaseExtractor):
    name = "docx"
    version = "1.0"
    supported_extensions = (".docx",)

    def __init__(self):
        super().__init__()
        try:
            import docx  # optional dependency check
        except ImportError as e:
            raise AIDocValueError(
                "DOCX extraction requires 'python-docx'. "
                "Install via: pip install aidoc-sdk[docx]"
            ) from e

    def extract(self, path: Path) -> RawDocument:
        import docx
        
        doc = docx.Document(path)

        elements: List = []
        order = 0

        # Extract block-level elements in order
        for block in doc.element.body:

            tag = block.tag.lower()

            # Paragraph
            if tag.endswith("}p"):
                paragraph = doc._element.xpath(
                    f'//w:p[@w14:paraId="{block.get(qn("w14:paraId"))}"]'
                )

            # Instead of fragile XML matching, iterate properly below

        # Safer iteration preserving order
        for child in doc.element.body:

            tag = child.tag.lower()

            # Paragraph
            if tag.endswith("}p"):
                paragraph = next(
                        p for p in doc.paragraphs
                        if p._element == child
                    )
                text = paragraph.text or ""

                style_name = paragraph.style.name if paragraph.style else ""

                if style_name.startswith("Heading"):
                    try:
                        level = int(style_name.replace("Heading", "").strip())
                    except Exception:
                        level = 1

                    elements.append(
                        RawHeading(
                            order=order,
                            text=text,
                            level=level,
                        )
                    )
                else:
                    elements.append(
                        RawParagraph(
                            order=order,
                            text=text,
                        )
                    )

                order += 1

            # Table
            elif tag.endswith("}tbl"):
                table = next(
                        t for t in doc.tables
                        if t._element == child
                    )

                rows: List[List[str]] = []

                for row in table.rows:
                    row_cells = [cell.text or "" for cell in row.cells]
                    rows.append(row_cells)

                elements.append(
                    RawTable(
                        order=order,
                        rows=rows,
                    )
                )

                order += 1

        return RawDocument(elements=elements)


# Extractor registered via decorator
